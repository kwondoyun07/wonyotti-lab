"""1단계: 공개 파일 구조 확인 & parquet 변환.

    python -m src.inspect_data inspect <파일>
    python -m src.inspect_data convert <파일> <출력.parquet> [--sheet 시트명]

지원 형식: .csv / .csv.gz / .xlsx
zip은 inspect로 안의 목록만 보여주니, 압축을 푼 뒤 각 파일로 다시 실행하세요.
변환 결과는 모든 값을 문자열로 저장합니다. 타입 변환은 다음 단계에서 명시적으로 합니다.
엑셀 시트 한 장은 약 104만 행까지라 140만 행은 여러 시트에 나뉘어 있을 수 있습니다.
convert는 첫 시트와 헤더가 같은 시트를 자동으로 이어 붙입니다.
"""

from __future__ import annotations

import argparse
import datetime as dt
import zipfile
from pathlib import Path

import pandas as pd
import pyarrow as pa
import pyarrow.parquet as pq

CHUNK = 200_000
ENCODINGS = ("utf-8-sig", "cp949")
HINTS = {
    "체결내역(executions)": {"execID", "symbol", "side", "lastQty", "lastPx", "execType", "ordType", "timestamp"},
    "지갑내역(wallet)": {"transactType", "amount", "walletBalance", "transactTime", "currency"},
}


def _is_csv(path: Path) -> bool:
    name = path.name.lower()
    return name.endswith(".csv") or name.endswith(".csv.gz")


def _csv_encoding(path: Path) -> str:
    for enc in ENCODINGS:
        try:
            pd.read_csv(path, nrows=200, encoding=enc, dtype=str)
            return enc
        except UnicodeDecodeError:
            continue
    raise SystemExit(f"인코딩을 알 수 없습니다: {path}")


def _unique_header(values) -> list[str]:
    seen: dict[str, int] = {}
    out = []
    for i, v in enumerate(values):
        name = str(v).strip() if v not in (None, "") else f"col_{i}"
        if name in seen:
            seen[name] += 1
            name = f"{name}_{seen[name]}"
        else:
            seen[name] = 0
        out.append(name)
    return out


def _to_str(v):
    if v is None:
        return None
    if isinstance(v, (dt.datetime, dt.date)):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def _frame(rows, cols) -> pd.DataFrame:
    n = len(cols)
    fixed = [[_to_str(v) for v in (list(r[:n]) + [None] * (n - len(r)))] for r in rows]
    return pd.DataFrame(fixed, columns=cols)


def _guess(columns) -> str:
    cols = set(columns)
    scores = {k: len(cols & v) for k, v in HINTS.items()}
    best = max(scores, key=scores.get)
    if scores[best] >= 3:
        return f"{best}로 보입니다 (BitMEX 기본 필드 {scores[best]}개 일치)"
    return "BitMEX 기본 필드명과 다릅니다 → config/columns.yaml에 직접 매핑하세요"


def _report(title: str, cols: list[str], head: pd.DataFrame, rows: str) -> None:
    print(f"\n[{title}] 행 수: {rows}")
    print("컬럼:", ", ".join(f"{i}:{c}" for i, c in enumerate(cols)))
    print("판단:", _guess(cols))
    with pd.option_context("display.width", 200, "display.max_columns", 40, "display.max_colwidth", 30):
        print(head.to_string(index=False))


def inspect(path: Path) -> None:
    if not path.exists():
        raise SystemExit(f"파일이 없습니다: {path}")
    print(f"# {path}  ({path.stat().st_size / 1e6:,.1f} MB)")

    if path.suffix.lower() == ".zip":
        with zipfile.ZipFile(path) as zf:
            print("zip 안의 파일 (압축을 푼 뒤 각 파일을 inspect 하세요):")
            for info in zf.infolist():
                print(f"  {info.filename}  {info.file_size / 1e6:,.1f} MB")
        return

    if _is_csv(path):
        enc = _csv_encoding(path)
        head = pd.read_csv(path, nrows=5, encoding=enc, dtype=str)
        n = sum(len(ch) for ch in pd.read_csv(path, encoding=enc, dtype=str, usecols=[0], chunksize=CHUNK))
        _report(f"csv, 인코딩 {enc}", list(head.columns), head, f"{n:,}")
        return

    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            for ws in wb.worksheets:
                it = ws.iter_rows(values_only=True)
                header = next(it, None)
                if header is None:
                    print(f"\n[시트 '{ws.title}'] 비어 있음")
                    continue
                cols = _unique_header(header)
                sample = [r for r in (next(it, None) for _ in range(5)) if r is not None]
                rows = f"약 {ws.max_row - 1:,} (시트 정보 기준)" if ws.max_row else "변환 시 확인"
                _report(f"시트 '{ws.title}'", cols, _frame(sample, cols), rows)
        finally:
            wb.close()
        return

    raise SystemExit(f"지원하지 않는 형식입니다: {path.suffix}")


def _convert_csv(path: Path, out: Path) -> int:
    enc = _csv_encoding(path)
    writer, schema, cols, total = None, None, None, 0
    try:
        for chunk in pd.read_csv(
            path, encoding=enc, dtype=str, chunksize=CHUNK, keep_default_na=False, na_values=[""]
        ):
            if writer is None:
                cols = _unique_header(chunk.columns)
                schema = pa.schema([(c, pa.string()) for c in cols])
                writer = pq.ParquetWriter(out, schema)
            chunk.columns = cols
            writer.write_table(pa.Table.from_pandas(chunk.astype("string"), schema=schema, preserve_index=False))
            total += len(chunk)
            print(f"  {total:,}행...", end="\r")
    finally:
        if writer is not None:
            writer.close()
    return total


def _write_rows(writer, schema, cols, rows) -> int:
    frame = _frame(rows, cols)
    writer.write_table(pa.Table.from_pandas(frame.astype("string"), schema=schema, preserve_index=False))
    return len(rows)


def _convert_xlsx(path: Path, out: Path, sheet: str | None) -> int:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    sheets = [wb[sheet]] if sheet else wb.worksheets
    writer, schema, cols, total = None, None, None, 0
    try:
        for ws in sheets:
            it = ws.iter_rows(values_only=True)
            header = next(it, None)
            if header is None:
                continue
            hcols = _unique_header(header)
            if cols is None:
                cols = hcols
                schema = pa.schema([(c, pa.string()) for c in cols])
                writer = pq.ParquetWriter(out, schema)
            elif hcols != cols:
                print(f"  [건너뜀] 시트 '{ws.title}'는 헤더가 다릅니다 → --sheet '{ws.title}'로 따로 변환하세요")
                continue
            batch = []
            for row in it:
                if row is None or all(v is None for v in row):
                    continue
                batch.append(row)
                if len(batch) >= CHUNK:
                    total += _write_rows(writer, schema, cols, batch)
                    batch = []
                    print(f"  {total:,}행...", end="\r")
            if batch:
                total += _write_rows(writer, schema, cols, batch)
            print(f"  시트 '{ws.title}' 완료 (누적 {total:,}행)")
    finally:
        wb.close()
        if writer is not None:
            writer.close()
    return total


def convert(path: Path, out: Path, sheet: str | None = None) -> None:
    if not path.exists():
        raise SystemExit(f"파일이 없습니다: {path}")
    out.parent.mkdir(parents=True, exist_ok=True)
    if _is_csv(path):
        total = _convert_csv(path, out)
    elif path.suffix.lower() in (".xlsx", ".xlsm"):
        total = _convert_xlsx(path, out, sheet)
    else:
        raise SystemExit(f"지원하지 않는 형식입니다: {path.suffix} (zip은 먼저 압축을 풀어주세요)")
    names = pq.read_schema(out).names
    print(f"\n완료: {out}  ({total:,}행, 컬럼 {len(names)}개)")
    print("컬럼:", ", ".join(names))
    print("판단:", _guess(names))


def main(argv=None) -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    p_inspect = sub.add_parser("inspect", help="파일 구조 확인")
    p_inspect.add_argument("path", type=Path)
    p_convert = sub.add_parser("convert", help="parquet으로 변환")
    p_convert.add_argument("path", type=Path)
    p_convert.add_argument("out", type=Path)
    p_convert.add_argument("--sheet", default=None, help="xlsx에서 특정 시트만 변환")
    args = ap.parse_args(argv)
    if args.cmd == "inspect":
        inspect(args.path)
    else:
        convert(args.path, args.out, args.sheet)


if __name__ == "__main__":
    main()
