import pandas as pd
import pyarrow.parquet as pq
from openpyxl import Workbook

from src.inspect_data import convert, inspect


def test_csv_roundtrip(tmp_path, capsys):
    src = tmp_path / "exec.csv"
    pd.DataFrame({"timestamp": ["2019-01-01T00:00:00Z", ""], "symbol": ["XBTUSD", "XBTUSD"],
                  "side": ["Buy", "Sell"], "lastQty": [100, 200], "lastPx": [4000.5, 4001],
                  "execType": ["Trade", "Trade"]}).to_csv(src, index=False)
    inspect(src)
    assert "체결내역" in capsys.readouterr().out
    out = tmp_path / "exec.parquet"
    convert(src, out)
    df = pd.read_parquet(out)
    assert len(df) == 2 and df["lastPx"].tolist() == ["4000.5", "4001.0"]
    assert pd.isna(df["timestamp"].iloc[1])


def test_xlsx_sheets_with_same_header_are_appended(tmp_path):
    src = tmp_path / "exec.xlsx"
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "part1"
    ws1.append(["timestamp", "symbol", "lastQty"])
    ws1.append(["2019-01-01T00:00:00Z", "XBTUSD", 100])
    ws2 = wb.create_sheet("part2")
    ws2.append(["timestamp", "symbol", "lastQty"])
    ws2.append(["2019-01-02T00:00:00Z", "XBTUSD", 200.0])
    ws3 = wb.create_sheet("wallet")
    ws3.append(["transactTime", "walletBalance"])
    ws3.append(["2019-01-01T00:00:00Z", 100000000])
    wb.save(src)

    out = tmp_path / "exec.parquet"
    convert(src, out)
    df = pd.read_parquet(out)
    assert df["lastQty"].tolist() == ["100", "200"]  # 헤더 다른 wallet 시트는 건너뜀

    out_w = tmp_path / "wallet.parquet"
    convert(src, out_w, sheet="wallet")
    assert pq.read_schema(out_w).names == ["transactTime", "walletBalance"]
